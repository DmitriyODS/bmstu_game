import os
import uuid
import io
import json
import re
import zipfile
import tempfile
import shutil
from flask import (render_template, redirect, url_for, request, flash,
                   jsonify, send_file, current_app)
from flask_login import current_user
from app.admin import admin_bp
from app.auth import admin_required, judge_required
from app.models import db, Quiz, Tour, Question, AnswerOption, MatchingItem, Team, TeamAnswer, GameState, User
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment


# ── Dashboard ──────────────────────────────────────────────────────────────────

@admin_bp.route('/')
@admin_required
def dashboard():
    quizzes = Quiz.query.order_by(Quiz.created_at.desc()).all()
    active_quiz = Quiz.query.filter_by(is_active=True).first()
    return render_template('admin/dashboard.html', quizzes=quizzes, active_quiz=active_quiz)


# ── Quiz list ──────────────────────────────────────────────────────────────────

@admin_bp.route('/quizzes')
@admin_required
def quizzes():
    all_quizzes = Quiz.query.order_by(Quiz.created_at.desc()).all()
    return render_template('admin/quizzes.html', quizzes=all_quizzes)


@admin_bp.route('/quizzes/new', methods=['GET', 'POST'])
@admin_required
def quiz_new():
    if request.method == 'POST':
        title = request.form.get('title', '').strip()
        description = request.form.get('description', '').strip()
        if not title:
            flash('Название обязательно', 'error')
            return render_template('admin/quiz_form.html', quiz=None)
        quiz = Quiz(title=title, description=description or None)
        db.session.add(quiz)
        db.session.flush()
        # Create default first tour
        tour = Tour(quiz_id=quiz.id, title='Тур 1', order=0)
        db.session.add(tour)
        db.session.commit()
        flash('Квиз создан', 'success')
        return redirect(url_for('admin.quiz_edit', quiz_id=quiz.id))
    return render_template('admin/quiz_form.html', quiz=None)


@admin_bp.route('/quizzes/<quiz_id>/edit', methods=['GET', 'POST'])
@admin_required
def quiz_edit(quiz_id):
    quiz = Quiz.query.get_or_404(quiz_id)
    if request.method == 'POST':
        quiz.title = request.form.get('title', quiz.title).strip()
        quiz.description = request.form.get('description', '').strip() or None
        db.session.commit()
        flash('Квиз обновлён', 'success')
    return render_template('admin/quiz_builder.html', quiz=quiz)


@admin_bp.route('/quizzes/<quiz_id>/upload-splash', methods=['POST', 'DELETE'])
@admin_required
def quiz_upload_splash(quiz_id):
    quiz = Quiz.query.get_or_404(quiz_id)
    if request.method == 'DELETE':
        quiz.splash_image = None
        db.session.commit()
        return jsonify({'ok': True})
    f = request.files.get('file')
    if not f:
        return jsonify({'error': 'No file'}), 400
    ext = (f.filename.rsplit('.', 1)[-1].lower() if f.filename and '.' in f.filename else '')
    if ext not in ('jpg', 'jpeg', 'png', 'gif', 'webp', 'avif', 'bmp', 'tiff', 'tif'):
        return jsonify({'error': 'Invalid file type'}), 400
    unique_name = f"{uuid.uuid4()}.{ext}"
    upload_dir = os.path.join(current_app.root_path, 'static', 'uploads', quiz.id, 'images')
    os.makedirs(upload_dir, exist_ok=True)
    f.save(os.path.join(upload_dir, unique_name))
    path = f'/static/uploads/{quiz.id}/images/{unique_name}'
    quiz.splash_image = path
    db.session.commit()
    return jsonify({'path': path})


@admin_bp.route('/quizzes/<quiz_id>/delete', methods=['POST'])
@admin_required
def quiz_delete(quiz_id):
    quiz = Quiz.query.get_or_404(quiz_id)
    quiz_id_copy = quiz.id
    from app.sockets.utils import broadcast
    broadcast('quiz_ended', {}, quiz_id_copy)
    db.session.delete(quiz)
    db.session.commit()
    flash('Квиз удалён', 'success')
    return redirect(url_for('admin.quizzes'))


@admin_bp.route('/quizzes/<quiz_id>/deactivate', methods=['POST'])
@admin_required
def quiz_deactivate(quiz_id):
    quiz = Quiz.query.get_or_404(quiz_id)
    quiz.is_active = False
    quiz.status = 'finished'
    db.session.commit()
    from app.sockets.utils import broadcast
    broadcast('quiz_ended', {}, quiz.id)
    flash('Квиз деактивирован', 'success')
    return redirect(url_for('admin.quizzes'))


@admin_bp.route('/quizzes/<quiz_id>/activate', methods=['POST'])
@admin_required
def quiz_activate(quiz_id):
    quiz = Quiz.query.get_or_404(quiz_id)
    # Deactivate all others
    Quiz.query.filter_by(is_active=True).update({'is_active': False})
    quiz.is_active = True
    quiz.status = 'active'
    db.session.flush()
    # Create or reset game state
    gs = GameState.query.filter_by(quiz_id=quiz.id).first()
    if not gs:
        gs = GameState(quiz_id=quiz.id, current_screen='splash', registration_open=False)
        db.session.add(gs)
    else:
        gs.current_screen = 'splash'
        gs.current_tour_id = None
        gs.current_question_id = None
        gs.timer_started_at = None
        gs.timer_seconds = None
        gs.registration_open = False
    db.session.commit()
    from app.sockets.utils import broadcast_all
    broadcast_all('quiz_activated', {'quiz_id': quiz_id})
    flash('Квиз активирован', 'success')
    return redirect(url_for('admin.quiz_control', quiz_id=quiz_id))


# ── Quiz Control ───────────────────────────────────────────────────────────────

@admin_bp.route('/quizzes/<quiz_id>/control')
@admin_required
def quiz_control(quiz_id):
    quiz = Quiz.query.get_or_404(quiz_id)
    gs = GameState.query.filter_by(quiz_id=quiz_id).first()
    if not gs:
        flash('Сначала активируйте квиз', 'warning')
        return redirect(url_for('admin.quiz_edit', quiz_id=quiz_id))
    return render_template('admin/quiz_control.html', quiz=quiz, gs=gs)


# ── Tours API ──────────────────────────────────────────────────────────────────

@admin_bp.route('/quizzes/<quiz_id>/tours', methods=['POST'])
@admin_required
def tour_create(quiz_id):
    quiz = Quiz.query.get_or_404(quiz_id)
    max_order = db.session.query(db.func.max(Tour.order)).filter_by(quiz_id=quiz_id).scalar()
    if max_order is None:
        max_order = -1
    count = Tour.query.filter_by(quiz_id=quiz_id).count()
    data = request.get_json(silent=True) or {}
    title = (data.get('title') or f'Тур {count + 1}').strip() or f'Тур {count + 1}'
    tour = Tour(quiz_id=quiz_id, title=title, order=max_order + 1)
    db.session.add(tour)
    db.session.commit()
    return jsonify({'id': tour.id, 'title': tour.title, 'order': tour.order})


@admin_bp.route('/tours/<tour_id>', methods=['GET', 'PUT', 'DELETE'])
@admin_required
def tour_update(tour_id):
    tour = Tour.query.get_or_404(tour_id)
    if request.method == 'DELETE':
        q_ids = [q.id for q in tour.questions]
        if q_ids:
            # Null out GameState references to questions/tour being deleted
            gs = GameState.query.filter_by(quiz_id=tour.quiz_id).first()
            if gs:
                if gs.current_tour_id == tour_id:
                    gs.current_tour_id = None
                if gs.current_question_id in q_ids:
                    gs.current_question_id = None
            # Delete TeamAnswers referencing these questions
            TeamAnswer.query.filter(TeamAnswer.question_id.in_(q_ids)).delete(synchronize_session=False)
        db.session.delete(tour)
        db.session.commit()
        return jsonify({'ok': True})
    if request.method == 'GET':
        return jsonify(_tour_dict(tour))
    data = request.get_json(silent=True) or {}
    for field in ('title', 'order'):
        if field in data:
            setattr(tour, field, data[field])
    db.session.commit()
    return jsonify(_tour_dict(tour))


@admin_bp.route('/tours/<tour_id>/reorder', methods=['POST'])
@admin_required
def tour_reorder(tour_id):
    data = request.get_json(silent=True) or []
    for item in data:
        if 'id' in item and 'order' in item:
            Tour.query.filter_by(id=item['id']).update({'order': item['order']})
    db.session.commit()
    return jsonify({'ok': True})


# Tour image upload
@admin_bp.route('/tours/<tour_id>/upload-splash', methods=['POST'])
@admin_required
def tour_upload_splash(tour_id):
    tour = Tour.query.get_or_404(tour_id)
    f = request.files.get('file')
    if not f:
        return jsonify({'error': 'No file'}), 400
    ext = (f.filename.rsplit('.', 1)[-1].lower() if f.filename and '.' in f.filename else '')
    if ext not in ('jpg', 'jpeg', 'png', 'gif', 'webp', 'avif', 'bmp', 'tiff', 'tif'):
        return jsonify({'error': 'Invalid file type'}), 400
    unique_name = f"{uuid.uuid4()}.{ext}"
    upload_dir = os.path.join(current_app.root_path, 'static', 'uploads',
                               tour.quiz_id, 'images')
    os.makedirs(upload_dir, exist_ok=True)
    f.save(os.path.join(upload_dir, unique_name))
    path = f'/static/uploads/{tour.quiz_id}/images/{unique_name}'
    tour.splash_image = path
    db.session.commit()
    return jsonify({'path': path})


# Quiz sound upload
@admin_bp.route('/quizzes/<quiz_id>/upload-sound', methods=['POST'])
@admin_required
def quiz_upload_sound(quiz_id):
    quiz = Quiz.query.get_or_404(quiz_id)
    sound_type = request.form.get('type', 'sound_start')  # sound_start | sound_mid | sound_end
    f = request.files.get('file')
    if not f:
        return jsonify({'error': 'No file'}), 400
    ext = (f.filename.rsplit('.', 1)[-1].lower() if f.filename and '.' in f.filename else '')
    if ext not in ('mp3', 'wav', 'ogg', 'm4a', 'aac', 'flac'):
        return jsonify({'error': 'Invalid audio type'}), 400
    unique_name = f"{uuid.uuid4()}.{ext}"
    upload_dir = os.path.join(current_app.root_path, 'static', 'uploads', quiz.id, 'audio')
    os.makedirs(upload_dir, exist_ok=True)
    f.save(os.path.join(upload_dir, unique_name))
    path = f'/static/uploads/{quiz.id}/audio/{unique_name}'
    field_map = {
        'sound_start': 'sound_start_path',
        'sound_mid': 'sound_mid_path',
        'sound_end': 'sound_end_path',
    }
    setattr(quiz, field_map.get(sound_type, 'sound_start_path'), path)
    db.session.commit()
    return jsonify({'path': path})


# Quiz settings update (sound_mid_seconds etc.)
@admin_bp.route('/quizzes/<quiz_id>/settings', methods=['PUT'])
@admin_required
def quiz_settings_update(quiz_id):
    quiz = Quiz.query.get_or_404(quiz_id)
    data = request.json or {}
    if 'sound_mid_seconds' in data:
        quiz.sound_mid_seconds = data['sound_mid_seconds']
    db.session.commit()
    return jsonify({'ok': True})


# ── Questions API ──────────────────────────────────────────────────────────────

@admin_bp.route('/tours/<tour_id>/questions', methods=['POST'])
@admin_required
def question_create(tour_id):
    tour = Tour.query.get_or_404(tour_id)
    max_order = db.session.query(db.func.max(Question.order)).filter_by(tour_id=tour_id).scalar()
    if max_order is None:
        max_order = -1
    data = request.get_json(silent=True) or {}
    q = Question(
        tour_id=tour_id,
        order=max_order + 1,
        question_type=data.get('question_type', 'text'),
        answer_type=data.get('answer_type', 'short_text'),
        text=data.get('text', ''),
        time_seconds=data.get('time_seconds', 60),
        points=data.get('points', 1),
    )
    db.session.add(q)
    db.session.commit()
    return jsonify(_question_dict(q))


@admin_bp.route('/questions/<question_id>', methods=['GET', 'PUT', 'DELETE'])
@admin_required
def question_update(question_id):
    q = Question.query.get_or_404(question_id)
    if request.method == 'DELETE':
        # Null out GameState reference if this is the current question
        gs = GameState.query.filter_by(current_question_id=question_id).first()
        if gs:
            gs.current_question_id = None
        # Delete TeamAnswers referencing this question
        TeamAnswer.query.filter_by(question_id=question_id).delete(synchronize_session=False)
        db.session.delete(q)
        db.session.commit()
        return jsonify({'ok': True})
    if request.method == 'GET':
        return jsonify(_question_dict(q))
    data = request.get_json(silent=True) or {}
    for field in ('question_type', 'answer_type', 'text', 'time_seconds', 'points',
                  'auto_check', 'correct_answer', 'order', 'audio_trim_start', 'audio_trim_end'):
        if field in data:
            setattr(q, field, data[field])
    # auto_check имеет смысл только для вариантов; принудительно нормализуем
    if q.answer_type not in ('single_choice', 'multiple_choice'):
        q.auto_check = False
    db.session.commit()
    return jsonify(_question_dict(q))


@admin_bp.route('/questions/<question_id>/reorder', methods=['POST'])
@admin_required
def question_reorder(question_id):
    data = request.get_json(silent=True) or []
    for item in data:
        if 'id' in item and 'order' in item:
            Question.query.filter_by(id=item['id']).update({'order': item['order']})
    db.session.commit()
    return jsonify({'ok': True})


@admin_bp.route('/questions/<question_id>/upload', methods=['POST'])
@admin_required
def question_upload(question_id):
    q = Question.query.get_or_404(question_id)
    file_type = request.form.get('type', 'image')  # image | audio | sound_start | sound_mid | sound_end
    f = request.files.get('file')
    if not f:
        return jsonify({'error': 'No file'}), 400

    ext = (f.filename.rsplit('.', 1)[-1].lower() if f.filename and '.' in f.filename else '')

    if file_type == 'image':
        if ext not in ('jpg', 'jpeg', 'png', 'gif', 'webp', 'avif', 'bmp', 'tiff', 'tif'):
            return jsonify({'error': 'Invalid image type'}), 400
        subdir = 'images'
    else:
        if ext not in ('mp3', 'wav', 'ogg', 'm4a', 'aac', 'flac'):
            return jsonify({'error': 'Invalid audio type'}), 400
        subdir = 'audio'

    tour = Tour.query.get(q.tour_id)
    upload_dir = os.path.join(current_app.root_path, 'static', 'uploads',
                               tour.quiz_id, subdir)
    os.makedirs(upload_dir, exist_ok=True)
    unique_name = f"{uuid.uuid4()}.{ext}"
    f.save(os.path.join(upload_dir, unique_name))
    path = f'/static/uploads/{tour.quiz_id}/{subdir}/{unique_name}'

    field_map = {
        'image': 'image_path',
        'audio': 'audio_path',
    }
    setattr(q, field_map.get(file_type, 'image_path'), path)
    if file_type == 'audio':
        q.audio_original_name = f.filename or unique_name
    db.session.commit()
    return jsonify({'path': path, 'original_name': q.audio_original_name})


@admin_bp.route('/questions/<question_id>/set-image-url', methods=['POST'])
@admin_required
def question_set_image_url(question_id):
    q = Question.query.get_or_404(question_id)
    url = (request.json or {}).get('url', '').strip()
    if not url:
        return jsonify({'error': 'No URL'}), 400
    q.image_path = url
    db.session.commit()
    return jsonify({'ok': True})


# ── Answer Options ─────────────────────────────────────────────────────────────

@admin_bp.route('/questions/<question_id>/options', methods=['GET', 'POST'])
@admin_required
def answer_options(question_id):
    q = Question.query.get_or_404(question_id)
    if request.method == 'GET':
        return jsonify([{'id': o.id, 'text': o.text, 'is_correct': o.is_correct, 'order': o.order}
                        for o in q.answer_options])
    data = request.get_json(silent=True) or {}
    max_order = db.session.query(db.func.max(AnswerOption.order)).filter_by(question_id=question_id).scalar()
    if max_order is None:
        max_order = -1
    opt = AnswerOption(
        question_id=question_id,
        text=data.get('text', ''),
        is_correct=data.get('is_correct', False),
        order=max_order + 1,
    )
    db.session.add(opt)
    db.session.commit()
    return jsonify({'id': opt.id, 'text': opt.text, 'is_correct': opt.is_correct, 'order': opt.order})


@admin_bp.route('/options/<option_id>', methods=['PUT', 'DELETE'])
@admin_required
def option_update(option_id):
    opt = AnswerOption.query.get_or_404(option_id)
    if request.method == 'DELETE':
        db.session.delete(opt)
        db.session.commit()
        return jsonify({'ok': True})
    data = request.get_json(silent=True) or {}
    if 'text' in data:
        opt.text = data['text']
    if 'is_correct' in data:
        opt.is_correct = bool(data['is_correct'])
    db.session.commit()
    return jsonify({'id': opt.id, 'text': opt.text, 'is_correct': opt.is_correct})


# ── Matching Items ─────────────────────────────────────────────────────────────

@admin_bp.route('/questions/<question_id>/matching', methods=['GET', 'POST'])
@admin_required
def matching_items(question_id):
    q = Question.query.get_or_404(question_id)
    if request.method == 'GET':
        return jsonify([{'id': m.id, 'left_text': m.left_text, 'right_text': m.right_text, 'order': m.order}
                        for m in q.matching_items])
    data = request.get_json(silent=True) or {}
    max_order = db.session.query(db.func.max(MatchingItem.order)).filter_by(question_id=question_id).scalar()
    if max_order is None:
        max_order = -1
    item = MatchingItem(
        question_id=question_id,
        left_text=data.get('left_text', ''),
        right_text=data.get('right_text', ''),
        order=max_order + 1,
    )
    db.session.add(item)
    db.session.commit()
    return jsonify({'id': item.id, 'left_text': item.left_text, 'right_text': item.right_text})


@admin_bp.route('/matching/<item_id>', methods=['PUT', 'DELETE'])
@admin_required
def matching_update(item_id):
    item = MatchingItem.query.get_or_404(item_id)
    if request.method == 'DELETE':
        db.session.delete(item)
        db.session.commit()
        return jsonify({'ok': True})
    data = request.get_json(silent=True) or {}
    if 'left_text' in data:
        item.left_text = data['left_text']
    if 'right_text' in data:
        item.right_text = data['right_text']
    db.session.commit()
    return jsonify({'id': item.id, 'left_text': item.left_text, 'right_text': item.right_text})


# ── Teams ──────────────────────────────────────────────────────────────────────

@admin_bp.route('/quizzes/<quiz_id>/teams')
@admin_required
def teams(quiz_id):
    quiz = Quiz.query.get_or_404(quiz_id)
    return render_template('admin/teams.html', quiz=quiz)


@admin_bp.route('/quizzes/<quiz_id>/teams/create', methods=['POST'])
@admin_required
def team_create(quiz_id):
    quiz = Quiz.query.get_or_404(quiz_id)
    name = request.form.get('name', '').strip()
    if not name:
        flash('Название обязательно', 'error')
        return redirect(url_for('admin.teams', quiz_id=quiz_id))
    token = str(uuid.uuid4())
    team = Team(quiz_id=quiz_id, name=name, session_token=token)
    db.session.add(team)
    db.session.commit()
    flash(f'Команда «{name}» создана', 'success')
    return redirect(url_for('admin.teams', quiz_id=quiz_id))


@admin_bp.route('/teams/<team_id>/delete', methods=['POST'])
@admin_required
def team_delete(team_id):
    team = Team.query.get_or_404(team_id)
    quiz_id = team.quiz_id
    db.session.delete(team)
    db.session.commit()
    flash('Команда удалена', 'success')
    return redirect(url_for('admin.teams', quiz_id=quiz_id))


@admin_bp.route('/teams/<team_id>/rename', methods=['POST'])
@admin_required
def team_rename(team_id):
    team = Team.query.get_or_404(team_id)
    name = request.form.get('name', '').strip()
    if name:
        team.name = name
        db.session.commit()
    return redirect(url_for('admin.teams', quiz_id=team.quiz_id))


# ── Settings ───────────────────────────────────────────────────────────────────

@admin_bp.route('/settings', methods=['GET', 'POST'])
@admin_required
def settings():
    judges = User.query.filter_by(role='judge').all()
    if request.method == 'POST':
        action = request.form.get('action')
        if action == 'change_credentials':
            new_username = request.form.get('username', '').strip()
            new_password = request.form.get('password', '').strip()
            if new_username:
                current_user.username = new_username
            if new_password:
                current_user.set_password(new_password)
            db.session.commit()
            flash('Данные обновлены', 'success')
        elif action == 'create_judge':
            username = request.form.get('judge_username', '').strip()
            password = request.form.get('judge_password', '').strip()
            if username and password:
                if User.query.filter_by(username=username).first():
                    flash('Пользователь уже существует', 'error')
                else:
                    judge = User(username=username, role='judge')
                    judge.set_password(password)
                    db.session.add(judge)
                    db.session.commit()
                    flash(f'Судья {username} создан', 'success')
        elif action == 'delete_judge':
            judge_id = request.form.get('judge_id')
            judge = User.query.get(judge_id)
            if judge and judge.role == 'judge':
                db.session.delete(judge)
                db.session.commit()
                flash('Судья удалён', 'success')
        return redirect(url_for('admin.settings'))
    return render_template('admin/settings.html', judges=judges)


# ── Export Excel ───────────────────────────────────────────────────────────────

@admin_bp.route('/export/<quiz_id>')
@admin_required
def export_excel(quiz_id):
    quiz = Quiz.query.get_or_404(quiz_id)
    teams_list = Team.query.filter_by(quiz_id=quiz_id).all()
    tours_list = quiz.tours

    wb = openpyxl.Workbook()

    # Sheet 1: Game results
    ws1 = wb.active
    ws1.title = 'Итоги игры'
    ws1.append(['Место', 'Команда', 'Баллы'])
    team_scores = []
    for team in teams_list:
        total = sum(a.score for a in team.answers if a.score)
        team_scores.append((team.name, total))
    team_scores.sort(key=lambda x: -x[1])
    for i, (name, score) in enumerate(team_scores, 1):
        ws1.append([i, name, score])

    # Sheet 2: By tours
    ws2 = wb.create_sheet('По турам')
    tour_headers = ['Команда'] + [t.title for t in tours_list]
    ws2.append(tour_headers)
    for team in teams_list:
        row = [team.name]
        for tour in tours_list:
            q_ids = [q.id for q in tour.questions]
            tour_score = sum(
                a.score for a in team.answers if a.question_id in q_ids and a.score
            )
            row.append(tour_score)
        ws2.append(row)

    # Sheet 3: By questions
    ws3 = wb.create_sheet('По вопросам')
    all_questions = []
    for tour in tours_list:
        all_questions.extend(tour.questions)
    q_headers = ['Команда'] + [f'Q{i+1}' for i in range(len(all_questions))]
    ws3.append(q_headers)
    for team in teams_list:
        row = [team.name]
        for q in all_questions:
            ans = next((a for a in team.answers if a.question_id == q.id), None)
            if ans:
                row.append(f"{ans.score} ({'✓' if ans.is_correct else '✗'})")
            else:
                row.append('')
        ws3.append(row)

    # Sheet 4: Answer details
    ws4 = wb.create_sheet('Детали ответов')
    ws4.append(['Команда', 'Вопрос', 'Ответ', 'Балл', 'Верно', 'Проверяющий', 'Время'])
    for team in teams_list:
        for ans in team.answers:
            q = ans.question
            checker_name = ans.checker.username if ans.checker else ''
            answer_text = ans.answer_text or str(ans.selected_options or ans.matching_pairs or '')
            ws4.append([
                team.name,
                q.text[:50] if q.text else '',
                answer_text[:100],
                ans.score,
                'Да' if ans.is_correct else ('Нет' if ans.is_correct is False else '?'),
                checker_name,
                str(ans.submitted_at),
            ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'quiz_{quiz.title}.xlsx',
    )


# ── JSON Export / Import ────────────────────────────────────────────────────────

@admin_bp.route('/quizzes/<quiz_id>/export-json')
@admin_required
def quiz_export_json(quiz_id):
    quiz = Quiz.query.get_or_404(quiz_id)
    data = {
        'title': quiz.title,
        'description': quiz.description,
        'tours': [],
    }
    for tour in quiz.tours:
        tour_data = {
            'title': tour.title,
            'order': tour.order,
            'questions': [],
        }
        for q in tour.questions:
            tour_data['questions'].append({
                'order': q.order,
                'question_type': q.question_type,
                'answer_type': q.answer_type,
                'text': q.text,
                'time_seconds': q.time_seconds,
                'points': q.points,
                'auto_check': q.auto_check,
                'correct_answer': q.correct_answer,
                'answer_options': [
                    {'text': o.text, 'is_correct': o.is_correct, 'order': o.order}
                    for o in q.answer_options
                ],
                'matching_items': [
                    {'left_text': m.left_text, 'right_text': m.right_text, 'order': m.order}
                    for m in q.matching_items
                ],
            })
        data['tours'].append(tour_data)
    buf = io.BytesIO(json.dumps(data, ensure_ascii=False, indent=2).encode('utf-8'))
    buf.seek(0)
    safe_title = re.sub(r'[\\/:*?"<>|]', '_', quiz.title).strip() or quiz_id
    return send_file(buf, mimetype='application/json', as_attachment=True,
                     download_name=f'{safe_title}.json')


@admin_bp.route('/quizzes/import-json', methods=['POST'])
@admin_required
def quiz_import_json():
    f = request.files.get('file')
    if not f or not f.filename.endswith('.json'):
        flash('Выберите файл .json', 'error')
        return redirect(url_for('admin.quizzes'))
    try:
        data = json.loads(f.read().decode('utf-8'))
    except Exception:
        flash('Не удалось прочитать файл', 'error')
        return redirect(url_for('admin.quizzes'))

    quiz = Quiz(
        title=data.get('title', 'Импортированный квиз'),
        description=data.get('description'),
    )
    db.session.add(quiz)
    db.session.flush()

    for tour_data in data.get('tours', []):
        tour = Tour(
            quiz_id=quiz.id,
            title=tour_data.get('title', 'Тур'),
            order=tour_data.get('order', 0),
        )
        db.session.add(tour)
        db.session.flush()
        for q_data in tour_data.get('questions', []):
            q = Question(
                tour_id=tour.id,
                order=q_data.get('order', 0),
                question_type=q_data.get('question_type', 'text'),
                answer_type=q_data.get('answer_type', 'short_text'),
                text=q_data.get('text'),
                image_path=q_data.get('image_path'),
                audio_path=q_data.get('audio_path'),
                time_seconds=q_data.get('time_seconds', 60),
                points=q_data.get('points', 1),
                auto_check=q_data.get('auto_check', False),
                correct_answer=q_data.get('correct_answer'),
            )
            db.session.add(q)
            db.session.flush()
            for o_data in q_data.get('answer_options', []):
                db.session.add(AnswerOption(
                    question_id=q.id,
                    text=o_data.get('text', ''),
                    is_correct=o_data.get('is_correct', False),
                    order=o_data.get('order', 0),
                ))
            for m_data in q_data.get('matching_items', []):
                db.session.add(MatchingItem(
                    question_id=q.id,
                    left_text=m_data.get('left_text', ''),
                    right_text=m_data.get('right_text', ''),
                    order=m_data.get('order', 0),
                ))

    db.session.commit()
    flash(f'Квиз «{quiz.title}» импортирован', 'success')
    return redirect(url_for('admin.quiz_edit', quiz_id=quiz.id))


# ── ZIP Export / Import ─────────────────────────────────────────────────────────

@admin_bp.route('/quizzes/<quiz_id>/export-zip')
@admin_required
def quiz_export_zip(quiz_id):
    quiz = Quiz.query.get_or_404(quiz_id)
    added_paths = set()

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:

        def add_media(path):
            if not path:
                return None
            full_path = os.path.join(current_app.root_path, path.lstrip('/'))
            if not os.path.exists(full_path):
                return None
            subdir = 'audio' if '/audio/' in path else 'images'
            filename = os.path.basename(full_path)
            zip_path = f'media/{subdir}/{filename}'
            if zip_path not in added_paths:
                zf.write(full_path, zip_path)
                added_paths.add(zip_path)
            return zip_path

        data = {
            'version': 1,
            'title': quiz.title,
            'description': quiz.description,
            'sound_mid_seconds': quiz.sound_mid_seconds,
            'splash_image': add_media(quiz.splash_image),
            'sound_start': add_media(quiz.sound_start_path),
            'sound_mid': add_media(quiz.sound_mid_path),
            'sound_end': add_media(quiz.sound_end_path),
            'tours': [],
        }

        for tour in quiz.tours:
            tour_data = {
                'title': tour.title,
                'order': tour.order,
                'splash_image': add_media(tour.splash_image),
                'questions': [],
            }
            for q in tour.questions:
                tour_data['questions'].append({
                    'order': q.order,
                    'question_type': q.question_type,
                    'answer_type': q.answer_type,
                    'text': q.text,
                    'time_seconds': q.time_seconds,
                    'points': q.points,
                    'auto_check': q.auto_check,
                    'correct_answer': q.correct_answer,
                    'image_path': add_media(q.image_path),
                    'audio_path': add_media(q.audio_path),
                    'audio_original_name': q.audio_original_name,
                    'audio_trim_start': q.audio_trim_start,
                    'audio_trim_end': q.audio_trim_end,
                    'answer_options': [
                        {'text': o.text, 'is_correct': o.is_correct, 'order': o.order}
                        for o in q.answer_options
                    ],
                    'matching_items': [
                        {'left_text': m.left_text, 'right_text': m.right_text, 'order': m.order}
                        for m in q.matching_items
                    ],
                })
            data['tours'].append(tour_data)

        zf.writestr('quiz.json', json.dumps(data, ensure_ascii=False, indent=2))

    buf.seek(0)
    safe_title = re.sub(r'[\\/:*?"<>|]', '_', quiz.title).strip() or quiz_id
    return send_file(buf, mimetype='application/zip', as_attachment=True,
                     download_name=f'{safe_title}.zip')


@admin_bp.route('/quizzes/import-zip', methods=['POST'])
@admin_required
def quiz_import_zip():
    f = request.files.get('file')
    if not f or not f.filename.lower().endswith('.zip'):
        flash('Выберите файл .zip', 'error')
        return redirect(url_for('admin.quizzes'))

    tmpdir = tempfile.mkdtemp()
    try:
        try:
            with zipfile.ZipFile(f, 'r') as zf:
                zf.extractall(tmpdir)
        except zipfile.BadZipFile:
            flash('Файл не является корректным ZIP-архивом', 'error')
            return redirect(url_for('admin.quizzes'))

        json_path = os.path.join(tmpdir, 'quiz.json')
        if not os.path.exists(json_path):
            flash('Неверный формат архива: отсутствует quiz.json', 'error')
            return redirect(url_for('admin.quizzes'))

        with open(json_path, 'r', encoding='utf-8') as jf:
            data = json.load(jf)

        quiz = Quiz(
            title=data.get('title', 'Импортированный квиз'),
            description=data.get('description'),
            sound_mid_seconds=data.get('sound_mid_seconds'),
        )
        db.session.add(quiz)
        db.session.flush()

        upload_base = os.path.join(current_app.root_path, 'static', 'uploads', quiz.id)

        def copy_media(zip_rel_path):
            if not zip_rel_path:
                return None
            src = os.path.join(tmpdir, *zip_rel_path.split('/'))
            if not os.path.exists(src):
                return None
            parts = zip_rel_path.split('/')
            subdir = parts[1] if len(parts) >= 2 and parts[1] in ('images', 'audio') else 'images'
            dst_dir = os.path.join(upload_base, subdir)
            os.makedirs(dst_dir, exist_ok=True)
            filename = os.path.basename(zip_rel_path)
            dst = os.path.join(dst_dir, filename)
            shutil.copy2(src, dst)
            return f'/static/uploads/{quiz.id}/{subdir}/{filename}'

        quiz.splash_image = copy_media(data.get('splash_image'))
        quiz.sound_start_path = copy_media(data.get('sound_start'))
        quiz.sound_mid_path = copy_media(data.get('sound_mid'))
        quiz.sound_end_path = copy_media(data.get('sound_end'))

        for tour_data in data.get('tours', []):
            tour = Tour(
                quiz_id=quiz.id,
                title=tour_data.get('title', 'Тур'),
                order=tour_data.get('order', 0),
            )
            db.session.add(tour)
            db.session.flush()
            tour.splash_image = copy_media(tour_data.get('splash_image'))

            for q_data in tour_data.get('questions', []):
                q = Question(
                    tour_id=tour.id,
                    order=q_data.get('order', 0),
                    question_type=q_data.get('question_type', 'text'),
                    answer_type=q_data.get('answer_type', 'short_text'),
                    text=q_data.get('text'),
                    time_seconds=q_data.get('time_seconds', 60),
                    points=q_data.get('points', 1),
                    auto_check=q_data.get('auto_check', False),
                    correct_answer=q_data.get('correct_answer'),
                    audio_original_name=q_data.get('audio_original_name'),
                    audio_trim_start=float(q_data.get('audio_trim_start') or 0),
                    audio_trim_end=q_data.get('audio_trim_end'),
                )
                db.session.add(q)
                db.session.flush()
                q.image_path = copy_media(q_data.get('image_path'))
                q.audio_path = copy_media(q_data.get('audio_path'))

                for o_data in q_data.get('answer_options', []):
                    db.session.add(AnswerOption(
                        question_id=q.id,
                        text=o_data.get('text', ''),
                        is_correct=o_data.get('is_correct', False),
                        order=o_data.get('order', 0),
                    ))
                for m_data in q_data.get('matching_items', []):
                    db.session.add(MatchingItem(
                        question_id=q.id,
                        left_text=m_data.get('left_text', ''),
                        right_text=m_data.get('right_text', ''),
                        order=m_data.get('order', 0),
                    ))

        db.session.commit()
        flash(f'Квиз «{quiz.title}» импортирован из ZIP', 'success')
        return redirect(url_for('admin.quiz_edit', quiz_id=quiz.id))

    except Exception as e:
        db.session.rollback()
        flash(f'Ошибка при импорте: {str(e)}', 'error')
        return redirect(url_for('admin.quizzes'))
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


# ── QR API ─────────────────────────────────────────────────────────────────────

@admin_bp.route('/api/qr/<quiz_id>')
def qr_code(quiz_id):
    import qrcode
    base_url = request.host_url.rstrip('/')
    url = f"{base_url}/play/{quiz_id}"
    img = qrcode.make(url)
    buf = io.BytesIO()
    img.save(buf, format='PNG')
    buf.seek(0)
    return send_file(buf, mimetype='image/png')


# ── Helpers ─────────────────────────────────────────────────────────────────────

def _tour_dict(tour):
    return {
        'id': tour.id,
        'title': tour.title,
        'order': tour.order,
        'splash_image': tour.splash_image,
    }


def _question_dict(q):
    return {
        'id': q.id,
        'tour_id': q.tour_id,
        'order': q.order,
        'question_type': q.question_type,
        'answer_type': q.answer_type,
        'text': q.text,
        'image_path': q.image_path,
        'audio_path': q.audio_path,
        'audio_original_name': q.audio_original_name,
        'audio_trim_start': q.audio_trim_start if q.audio_trim_start is not None else 0.0,
        'audio_trim_end': q.audio_trim_end,
        'time_seconds': q.time_seconds,
        'points': q.points,
        'auto_check': q.auto_check,
        'correct_answer': q.correct_answer,
        'answer_options': [
            {'id': o.id, 'text': o.text, 'is_correct': o.is_correct, 'order': o.order}
            for o in q.answer_options
        ],
        'matching_items': [
            {'id': m.id, 'left_text': m.left_text, 'right_text': m.right_text, 'order': m.order}
            for m in q.matching_items
        ],
    }
